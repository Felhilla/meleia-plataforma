"""Genera los Excel de muestra para las pruebas de contrato de la carga por Excel.
Datos 100 % inventados. La estructura imita el Excel de mapeo de GH (Módulo I) y la plantilla del Módulo IV.
Uso: python3 scripts/generar-fixtures-excel.py  (requiere openpyxl)"""
import json, pathlib
from openpyxl import Workbook

RAIZ = pathlib.Path(__file__).resolve().parent.parent
SALIDA = RAIZ / 'tests' / 'fixtures'
ESCALA_15 = 'Muy Baja(1); Baja(2); Media(3); Alta(4); Muy Alta(5)'

def mapeo_gh():
    wb = Workbook(); wb.active.title = 'Índice'; wb.active['A1'] = 'Índice (se ignora)'
    wb.create_sheet('Metodología')['A1'] = 'Metodología (se ignora)'
    # Hoja comparativa: se ignora (empieza por «Comp.»).
    comp = wb.create_sheet('Comp. Ins. Gub'); comp.append([]); comp['B13'], comp['C13'], comp['D13'] = 'Institución', 'Influencia', 'Impacto'

    def hoja(titulo, fila_encabezado, columnas, filas):
        ws = wb.create_sheet(titulo)
        ws.cell(fila_encabezado - 2, 2, 'Texto introductorio de la hoja (se ignora)')
        for i, (cab, escala) in enumerate(columnas, start=2):
            if cab: ws.cell(fila_encabezado, i, cab)
            if escala: ws.cell(fila_encabezado + 1, i, escala)
        for col in ('B', 'C', 'F'):
            ws.merge_cells(f'{col}{fila_encabezado}:{col}{fila_encabezado + 1}')
        for j, fila in enumerate(filas, start=fila_encabezado + 2):
            for i, v in enumerate(fila, start=2): ws.cell(j, i, v)
        return ws

    # Variante A: con Cargo, Contacto en una columna, Influencia e Impacto en texto + número (encabezado combinado).
    cols_a = [('Institución', None), ('Nombre', None), ('Cargo', None), ('Contacto', None), ('Caracterización', None), ('Posición', None),
              ('Dependencia', ESCALA_15), ('Tensión', ESCALA_15), ('Responsabilidad', ESCALA_15),
              ('Influencia', 'Muy Baja; Baja; Media; Alta; Muy Alta'), (None, '(1-5)'), ('Impacto', 'Muy Bajo; Bajo; Medio; Alto; Muy Alto'), (None, '(1-5)'),
              ('Nivel de relacionamiento', None)]
    ws = hoja('Com. Aso', 12, cols_a, [
        ['JAC vereda El Ejemplo', 'Persona Inventada Uno', 'Presidente JAC', '3000000000', 'Junta que representa a la vereda vecina del proyecto.', 'Desfavorable', 'Media(3)', 'Alta(4)', 'Baja(2)', 'Media', 3, 'Muy Alto', 5, 'Gestión Cautelosa'],
        ['Asociación de Productores Ficticios', 'Persona Inventada Dos', 'Representante legal', 'correo@ejemplo.test', 'Asociación de productores de la zona.', 'Favorable', 'Baja(2)', 'Muy Baja(1)', 'Media(3)', 'baja', 2, 'Bajo', 2, 'Seguimiento periférico'],
        # Fila con error: influencia no reconocida.
        ['Colectivo Sin Nombre', None, 'Vocera', None, None, 'Neutral', 'Media(3)', 'Baja(2)', 'Baja(2)', 'Alto medio', None, 'Medio', 3, None],
    ])
    ws.merge_cells('K12:L12'); ws.merge_cells('M12:N12')

    # Variante B: sin Cargo y con Contacto dividido en Correo / Número (como algunas hojas de GH).
    cols_b = [('Institución', None), ('Nombre', None), ('Contacto', 'Correo'), (None, 'Número'), ('Caracterización', None), ('Posición', None),
              ('Dependencia', ESCALA_15), ('Tensión', ESCALA_15), ('Responsabilidad', ESCALA_15),
              ('Influencia', 'Muy Baja; Baja; Media; Alta; Muy Alta'), (None, '(1-5)'), ('Impacto', 'Muy Bajo; Bajo; Medio; Alto; Muy Alto'), (None, '(1-5)'),
              ('Nivel de relacionamiento', None)]
    ws = hoja('Ins. Gub', 12, cols_b, [
        ['Alcaldía Municipal de Ejemplo', 'Persona Inventada Tres', 'alcalde@ejemplo.test', '3000000001', 'Autoridad municipal del área de influencia.', 'Neutral', 'Media(3)', 'Baja(2)', 'Muy Alta(5)', 'Muy Alta', 5, 'Alto', 4, 'Involucramiento Estratégico'],
        # Nivel del archivo un paso arriba del calculado (Seguimiento periférico → Gestión activa): entra como ajuste con motivo 5.
        ['Gobernación Ficticia', 'Persona Inventada Cuatro', None, None, 'Autoridad departamental.', 'Favorable', 'Baja(2)', 'Muy Baja(1)', 'Alta(4)', 'Media', 3, 'Bajo', 2, 'Gestión Activa'],
    ])
    ws.merge_cells('D12:E12'); ws.merge_cells('K12:L12'); ws.merge_cells('M12:N12')
    wb.save(SALIDA / 'mapeo-gh-muestra.xlsx')

def plantilla_iv():
    hojas = [('Riesgos', 'riesgos'), ('Calificación técnica', 'riesgos-calificacion'), ('Grupos', 'riesgos-grupos'), ('Validaciones', 'riesgos-validaciones'), ('Medidas', 'riesgos-medidas')]
    derechos = list(json.loads((RAIZ / 'public/data/derechos.json').read_text())['derechos'])
    wb = Workbook(); wb.remove(wb.active)
    filas = {
        'riesgos': [{'codigo': 'R-01', 'nombre': 'Afectación del agua de la vereda', 'derecho_humano': derechos[0], 'ambitos': 'Operación propia; Comunidades del área de influencia'},
                    {'codigo': 'r-2', 'nombre': 'Código mal escrito (error esperado)', 'derecho_humano': derechos[0]}],
        'riesgos-calificacion': [{'riesgo': 'R-01', 'escala': 2, 'alcance': 3, 'irremediabilidad': 2, 'probabilidad': 'Media'}],
        'riesgos-grupos': [{'nombre': 'Mesa comunitaria inventada', 'categoria': 'Comunidades y asociaciones', 'participantes': 12, 'fecha': '2026-10-01'}],
        'riesgos-validaciones': [{'riesgo': 'R-01', 'grupo': 'Mesa comunitaria inventada'}],
        'riesgos-medidas': [{'codigo': 'A-01', 'titulo': 'Monitoreo participativo del agua', 'riesgos': 'R-01'}],
    }
    for titulo, coleccion in hojas:
        esquema = json.loads((RAIZ / f'public/config/esquemas/{coleccion}.json').read_text())
        ws = wb.create_sheet(titulo); campos = list(esquema['campos'])
        ws.append([esquema['campos'][c]['etiqueta'] for c in campos])
        for f in filas[coleccion]: ws.append([f.get(c) for c in campos])
    wb.save(SALIDA / 'plantilla-iv-muestra.xlsx')

if __name__ == '__main__':
    SALIDA.mkdir(parents=True, exist_ok=True); mapeo_gh(); plantilla_iv(); print('OK', SALIDA)
