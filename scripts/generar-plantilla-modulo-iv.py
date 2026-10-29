"""Genera la plantilla descargable de carga por Excel del Módulo IV.

Uso: python3 scripts/generar-plantilla-modulo-iv.py  (requiere openpyxl)
"""
import json
import pathlib

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


RAIZ = pathlib.Path(__file__).resolve().parent.parent
SALIDA = RAIZ / 'public' / 'plantillas' / 'plantilla-modulo-iv.xlsx'


def cargar_json(ruta):
    return json.loads(ruta.read_text(encoding='utf-8'))


def agregar_lista(ws_listas, valores):
    """Escribe una lista en la siguiente columna y devuelve su rango Excel."""
    columna = 1 if ws_listas['A1'].value is None else ws_listas.max_column + 1
    for fila, valor in enumerate(valores, start=1):
        ws_listas.cell(fila, columna, valor)
    letra = get_column_letter(columna)
    return f'Listas!${letra}$1:${letra}${len(valores)}'


def agregar_validacion(ws, columna, rango_lista):
    validacion = DataValidation(type='list', formula1=rango_lista, allow_blank=True)
    ws.add_data_validation(validacion)
    validacion.add(f'{get_column_letter(columna)}2:{get_column_letter(columna)}500')


def main():
    carga = cargar_json(RAIZ / 'public' / 'config' / 'kit-carga-excel.json')
    derechos = list(cargar_json(RAIZ / 'public' / 'data' / 'derechos.json')['derechos'])
    componentes = [
        componente['nombre']
        for componente in cargar_json(
            RAIZ / 'public' / 'config' / 'kit-riesgos-config.json'
        )['componentes']
    ]

    listas = {
        'derechos': derechos,
        'numeros_1_a_3': [1, 2, 3],
        'probabilidades': ['Alta', 'Media', 'Baja'],
        'tipos_impacto': ['Potencial', 'Manifestado'],
        'vinculaciones': ['Causa', 'Contribuye', 'Directamente vinculada'],
        'categorias_actor': [
            'Instituciones gubernamentales',
            'Comunidades y asociaciones',
            'Academia',
            'Empresas y gremios',
            'Medios de comunicación',
            'Empleados, contratistas y proveedores',
        ],
        'resultados': ['Confirma', 'Propone ajuste'],
        'componentes': componentes,
        'estados': ['Pendiente', 'En curso', 'Cumplida'],
    }

    validaciones = {
        'Riesgos': {'Derecho humano afectado': 'derechos'},
        'Calificación técnica': {
            'Escala': 'numeros_1_a_3',
            'Alcance': 'numeros_1_a_3',
            'Irremediabilidad': 'numeros_1_a_3',
            'Probabilidad': 'probabilidades',
            'Tipo de impacto': 'tipos_impacto',
            'Vinculación': 'vinculaciones',
        },
        'Grupos': {'Categoría de actor': 'categorias_actor'},
        'Validaciones': {
            'Resultado': 'resultados',
            'Escala propuesta': 'numeros_1_a_3',
            'Alcance propuesta': 'numeros_1_a_3',
            'Irremediabilidad propuesta': 'numeros_1_a_3',
            'Probabilidad propuesta': 'probabilidades',
            'Tipo de impacto propuesta': 'tipos_impacto',
            'Vinculación propuesta': 'vinculaciones',
        },
        'Medidas': {
            'Componente': 'componentes',
            'Estado': 'estados',
        },
    }

    wb = Workbook()
    wb.remove(wb.active)
    hojas = {}
    for hoja in carga['riesgos']['hojas']:
        ws = wb.create_sheet(hoja['hoja'])
        encabezados = [columna['encabezado'] for columna in hoja['columnas']]
        ws.append(encabezados)
        hojas[hoja['hoja']] = (ws, encabezados)

    ws_listas = wb.create_sheet('Listas')
    rangos_listas = {
        nombre: agregar_lista(ws_listas, valores)
        for nombre, valores in listas.items()
    }

    for nombre_hoja, (ws, encabezados) in hojas.items():
        for indice, encabezado in enumerate(encabezados, start=1):
            lista = validaciones.get(nombre_hoja, {}).get(encabezado)
            if lista:
                agregar_validacion(ws, indice, rangos_listas[lista])

    ws_listas.sheet_state = 'hidden'
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print('OK', SALIDA)


if __name__ == '__main__':
    main()
