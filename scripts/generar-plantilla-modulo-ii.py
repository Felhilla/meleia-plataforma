"""Genera la plantilla descargable de carga por Excel del Módulo II (debida diligencia en DDHH).

Uso: python3 scripts/generar-plantilla-modulo-ii.py  (requiere openpyxl)

La plantilla trae ya escritos los 159 indicadores y las 18 dimensiones del catálogo. El responsable
solo llena lo que se califica a mano; la plataforma calcula los promedios (KitEstudios.calcular).
"""
import json
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


RAIZ = pathlib.Path(__file__).resolve().parent.parent
SALIDA = RAIZ / 'public' / 'plantillas' / 'plantilla-modulo-ii.xlsx'
FILAS_EDITABLES = 'FFF7E6'
ENCABEZADO = 'E8EEF4'

INSTRUCCIONES = [
    ('Plantilla de carga · Módulo II · Debida diligencia en DDHH', True),
    ('', False),
    ('Cómo se llena', True),
    ('1. Hoja «Evaluación»: escribe el nombre y la fecha (AAAA-MM-DD) en la fila 2. El mes identifica la evaluación: '
     'si subes otro archivo con el mismo año y mes, reemplaza los datos de esa evaluación.', False),
    ('2. Hoja «Indicadores»: llena las columnas en color (Incorporado, Documentos analizados, Hallazgo, Calificación '
     'y Brecha y recomendación). No cambies ni borres las columnas Código, Etapa o tema, Criterio y Pregunta, '
     'ni borres filas.', False),
    ('3. Hoja «Hallazgos»: escribe los hallazgos y las brechas de cada dimensión, uno por línea '
     '(Alt + Enter para cambiar de línea dentro de la celda).', False),
    ('', False),
    ('Regla importante: la calificación en blanco vale 0', True),
    ('Una calificación vacía cuenta como 0 («No hay información») en todos los promedios. Si un indicador tiene '
     'evidencia, califícalo; si no la tiene, déjalo en 0. La vista previa dice cuántos indicadores quedaron en blanco '
     'antes de que confirmes la carga.', False),
    ('', False),
    ('Qué calcula la plataforma', True),
    ('No escribas promedios. La plataforma calcula cada criterio como el promedio de sus indicadores, cada etapa de la '
     'Guía OCDE como el promedio de sus criterios y cada tema de DDHH como el promedio de sus indicadores.', False),
    ('', False),
    ('Escala de calificación', True),
]


def cargar_json(ruta):
    return json.loads(ruta.read_text(encoding='utf-8'))


def encabezados(ws, titulos, anchos):
    ws.append(titulos)
    for indice, (celda, ancho) in enumerate(zip(ws[1], anchos), start=1):
        celda.font = Font(bold=True)
        celda.fill = PatternFill('solid', fgColor=ENCABEZADO)
        celda.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(indice)].width = ancho
    ws.freeze_panes = 'B2'


def marcar_editables(ws, columnas, filas):
    """Desbloquea y colorea las celdas que llena el responsable; el resto queda protegido."""
    for fila in range(2, filas + 2):
        for columna in range(1, ws.max_column + 1):
            celda = ws.cell(fila, columna)
            celda.alignment = Alignment(wrap_text=True, vertical='top')
            if columna in columnas:
                celda.protection = Protection(locked=False)
                celda.fill = PatternFill('solid', fgColor=FILAS_EDITABLES)
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def main():
    carga = cargar_json(RAIZ / 'public' / 'config' / 'kit-carga-excel.json')['ddhh']
    cat = cargar_json(RAIZ / 'public' / 'config' / 'kit-ddhh-catalogo.json')
    esquema_indicadores = cargar_json(RAIZ / 'public' / 'config' / 'esquemas' / 'ddhh-indicadores.json')
    hojas = {hoja['coleccion']: hoja for hoja in carga['hojas']}
    titulo = lambda coleccion, campo: next(c['encabezado'] for c in hojas[coleccion]['columnas'] if c['campo'] == campo)
    ejes = {eje['id']: eje for eje in cat['ejes']}
    componentes = {c['id']: c['nombre'] for eje in cat['ejes'] for c in eje['componentes']}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Instrucciones'
    ws.column_dimensions['A'].width = 110
    for texto, negrita in INSTRUCCIONES:
        ws.append([texto])
        ws.cell(ws.max_row, 1).font = Font(bold=negrita, size=13 if ws.max_row == 1 else 11)
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
    for nivel in cat['estandares'][0]['escala']['niveles']:
        ws.append([f"{nivel['valor']}: {nivel['descripcion']}"])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')

    # Evaluación: una fila.
    ws = wb.create_sheet(hojas['ddhh-evaluaciones']['hoja'])
    encabezados(ws, [c['encabezado'] for c in hojas['ddhh-evaluaciones']['columnas']], [50, 18])
    ws.append(['', ''])
    ws['B2'].number_format = 'yyyy-mm-dd'
    marcar_editables(ws, {1, 2}, 1)

    # Indicadores: las 159 filas del catálogo.
    hoja = hojas['ddhh-indicadores']
    columnas = [titulo('ddhh-indicadores', 'indicador'), *hoja['informativas'],
                *[c['encabezado'] for c in hoja['columnas'] if c['campo'] != 'indicador']]
    ws = wb.create_sheet(hoja['hoja'])
    encabezados(ws, columnas, [18, 30, 30, 60, 14, 40, 50, 14, 50])
    for indicador in cat['indicadores']:
        fila = {columnas[0]: indicador['id'], 'Etapa o tema': ejes[indicador['eje']]['nombre'],
                'Criterio': componentes.get(indicador.get('componente'), '—'), 'Pregunta': indicador['pregunta']}
        ws.append([fila.get(columna, '') for columna in columnas])
    editables = {columnas.index(c['encabezado']) + 1 for c in hoja['columnas'] if c['campo'] != 'indicador'}
    marcar_editables(ws, editables, len(cat['indicadores']))
    ultima = len(cat['indicadores']) + 1

    listas = wb.create_sheet('Listas')
    opciones = {'si': 'Sí', 'no': 'No', 'parcial': 'Parcial'}
    for fila, valor in enumerate(esquema_indicadores['campos']['incorporado']['valores'], start=1):
        listas.cell(fila, 1, opciones.get(valor, valor))
    listas.sheet_state = 'hidden'
    columna = get_column_letter(columnas.index(titulo('ddhh-indicadores', 'incorporado')) + 1)
    lista = DataValidation(type='list', formula1=f'Listas!$A$1:$A${len(opciones)}', allow_blank=True)
    ws.add_data_validation(lista)
    lista.add(f'{columna}2:{columna}{ultima}')
    columna = get_column_letter(columnas.index(titulo('ddhh-indicadores', 'calificacion')) + 1)
    escala = cat['estandares'][0]['escala']
    numero = DataValidation(type='decimal', operator='between', formula1=str(escala['min']), formula2=str(escala['max']),
                            allow_blank=True, showErrorMessage=True, errorTitle='Calificación',
                            error=f"Escribe un número de {escala['min']} a {escala['max']}.")
    ws.add_data_validation(numero)
    numero.add(f'{columna}2:{columna}{ultima}')

    # Hallazgos: las 18 dimensiones (dos generales y una por eje).
    hoja = hojas['ddhh-dimensiones']
    columnas = [titulo('ddhh-dimensiones', 'dimension'), *hoja['informativas'],
                *[c['encabezado'] for c in hoja['columnas'] if c['campo'] != 'dimension']]
    ws = wb.create_sheet(hoja['hoja'])
    encabezados(ws, columnas, [18, 45, 60, 60])
    dimensiones = [(d['id'], d['titulo']) for d in cat['dimensiones']['grandes']] + [(e['id'], e['nombre']) for e in cat['ejes']]
    for codigo, nombre in dimensiones:
        ws.append([codigo, nombre, '', ''])
    marcar_editables(ws, {3, 4}, len(dimensiones))

    wb.move_sheet('Listas', offset=len(wb.sheetnames))
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print('OK', SALIDA)


if __name__ == '__main__':
    main()
